from __future__ import annotations

from datetime import datetime
from pathlib import Path
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "UK Support Housing - Resellable Platform Package.docx"

SOURCES = [
    ROOT / "docs" / "RESELLABLE_PLATFORM_PACKAGE.md",
    ROOT / "docs" / "CRM_SYSTEM_OVERVIEW.md",
    ROOT / "docs" / "SECURITY.md",
    ROOT / "docs" / "OPERATIONS.md",
    ROOT / "docs" / "TROUBLESHOOTING.md",
]


def shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text: str, bold: bool = False) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.size = Pt(9)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_note(doc: Document, title: str, body: str, fill: str = "F8FAFC") -> None:
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    cell = table.cell(0, 0)
    shade_cell(cell, fill)
    p = cell.paragraphs[0]
    run = p.add_run(title)
    run.bold = True
    run.font.color.rgb = RGBColor(23, 32, 51)
    p.add_run(f"\n{body}")
    for paragraph in cell.paragraphs:
      paragraph.paragraph_format.space_after = Pt(3)


def add_markdown_table(doc: Document, rows: list[str]) -> None:
    parsed = [[cell.strip() for cell in row.strip().strip("|").split("|")] for row in rows if row.strip().startswith("|")]
    parsed = [row for row in parsed if not all(set(cell) <= {"-", ":"} for cell in row)]
    if not parsed:
        return
    table = doc.add_table(rows=len(parsed), cols=max(len(row) for row in parsed))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for r, row in enumerate(parsed):
        for c in range(len(table.columns)):
            text = row[c] if c < len(row) else ""
            set_cell_text(table.cell(r, c), text, bold=(r == 0))
            if r == 0:
                shade_cell(table.cell(r, c), "172033")
                for paragraph in table.cell(r, c).paragraphs:
                    for run in paragraph.runs:
                        run.font.color.rgb = RGBColor(255, 255, 255)


def add_markdown(doc: Document, text: str) -> None:
    lines = text.splitlines()
    code_mode = False
    code_lines: list[str] = []
    table_lines: list[str] = []

    def flush_table() -> None:
        nonlocal table_lines
        if table_lines:
            add_markdown_table(doc, table_lines)
            table_lines = []

    def flush_code() -> None:
        nonlocal code_lines
        if code_lines:
            p = doc.add_paragraph()
            p.style = "Code"
            p.add_run("\n".join(code_lines))
            code_lines = []

    for raw in lines:
        line = raw.rstrip()
        if line.startswith("```"):
            flush_table()
            if code_mode:
                flush_code()
                code_mode = False
            else:
                code_mode = True
            continue

        if code_mode:
            code_lines.append(line)
            continue

        if line.strip().startswith("|"):
            table_lines.append(line)
            continue
        flush_table()

        if not line.strip():
            continue

        heading = re.match(r"^(#{1,4})\s+(.*)", line)
        if heading:
            level = min(len(heading.group(1)), 3)
            doc.add_heading(heading.group(2), level=level)
            continue

        if line.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            p.add_run(line[2:])
            continue

        numbered = re.match(r"^\d+\.\s+(.*)", line)
        if numbered:
            p = doc.add_paragraph(style="List Number")
            p.add_run(numbered.group(1))
            continue

        p = doc.add_paragraph()
        p.add_run(line)

    flush_table()
    flush_code()


def build() -> None:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.72)
    section.right_margin = Inches(0.72)

    styles = doc.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(10.5)
    styles["Normal"].paragraph_format.space_after = Pt(5)
    styles["Title"].font.name = "Aptos Display"
    styles["Title"].font.size = Pt(28)
    styles["Title"].font.bold = True
    for name, size, color in [("Heading 1", 18, "172033"), ("Heading 2", 14, "243B53"), ("Heading 3", 11.5, "334E68")]:
        styles[name].font.name = "Aptos Display"
        styles[name].font.size = Pt(size)
        styles[name].font.bold = True
        styles[name].font.color.rgb = RGBColor.from_string(color)
        styles[name].paragraph_format.space_before = Pt(12)
        styles[name].paragraph_format.space_after = Pt(5)
    code = styles.add_style("Code", 1)
    code.font.name = "Consolas"
    code.font.size = Pt(8.5)
    code.paragraph_format.left_indent = Inches(0.2)
    code.paragraph_format.space_before = Pt(4)
    code.paragraph_format.space_after = Pt(8)

    title = doc.add_paragraph()
    title.style = "Title"
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("Resellable Platform Package\n")
    sub = title.add_run("Supported Housing PMS + CRM + Compliance Platform")
    sub.font.size = Pt(14)
    sub.font.bold = False
    sub.font.color.rgb = RGBColor(82, 96, 117)

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.add_run(f"Generated: {datetime.now().strftime('%d %B %Y %H:%M')}").italic = True

    add_note(
        doc,
        "Purpose",
        "This document packages the technical setup, white-label configuration, deployment, operations, security, CRM, and troubleshooting guidance into a single MS Word handover manual.",
        "FFF7D6",
    )

    doc.add_page_break()
    doc.add_heading("Document Sections", level=1)
    for item in [
        "Executive summary and white-label readiness",
        "Fresh deployment and new company setup",
        "Configuration reference and module map",
        "Codebase and database maps",
        "Operations, security and troubleshooting",
        "CRM technical overview",
        "Change log and verification notes",
    ]:
        doc.add_paragraph(item, style="List Bullet")

    for source in SOURCES:
        if source.exists():
            doc.add_page_break()
            add_markdown(doc, source.read_text(encoding="utf-8"))

    doc.add_page_break()
    doc.add_heading("Server Docker Verification", level=1)
    add_note(
        doc,
        "Completed on Ubuntu server",
        "Docker Engine and Docker Compose were installed. A separate Compose project named realtyos-test was built and started on app port 3001 and PostgreSQL port 5433. First-run setup completed, homepage returned 200, admin login returned 200, and unauthenticated ERP API returned 401.",
        "E8F7EE",
    )
    doc.add_paragraph("The live systemd deployment remains active on the production route. The Docker verification stack is separate from the live service.")

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
